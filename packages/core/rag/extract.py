"""Raw upload bytes -> plain text, by the file's suffix.

The upload route validated the suffix and the compressed size; this decides
the reader and bounds the work, because a few kilobytes of PDF or DOCX can
inflate to hundreds of megabytes and a page tree can be arbitrarily deep.
Text formats are UTF-8; PDF and DOCX are parsed for their text layer only
(no OCR), so a scanned PDF yields nothing and fails with a message the user
can act on. Every failure here is a PermanentIngestError: parsing is pure,
so a file that cannot be read now will not be readable on retry.
"""

import io
import zipfile

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError

from packages.core.rag.errors import PermanentIngestError

# Budgets, all static. A DOCX is a zip: its members' declared sizes are read
# before anything is inflated. PDF memory per page is pypdf's own, which these
# do not bound — the worker's memory fence and the redelivery stop in
# apps/worker/tasks.py are the backstop for a page that explodes.
MAX_INFLATED_BYTES = 25 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_TEXT_CHARS = 2_000_000

TOO_LONG = "the document is too long; split it and upload the parts"


def extract_text(raw: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _pdf(raw)
    elif name.endswith(".docx"):
        text = _docx(raw)
    else:
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise PermanentIngestError("the file is not valid UTF-8 text") from exc
    if len(text) > MAX_TEXT_CHARS:
        raise PermanentIngestError(TOO_LONG)
    return text


def _pdf(raw: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(raw))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise PermanentIngestError(TOO_LONG)
        pages: list[str] = []
        total = 0
        for page in reader.pages:
            text = (page.extract_text() or "").strip()
            total += len(text)
            if total > MAX_TEXT_CHARS:
                raise PermanentIngestError(TOO_LONG)
            if text:
                pages.append(text)
    except PermanentIngestError:
        raise
    except FileNotDecryptedError as exc:
        raise PermanentIngestError(
            "the PDF is password-protected; remove the password and upload again"
        ) from exc
    except Exception as exc:  # noqa: BLE001 — pypdf raises many types; all permanent
        raise PermanentIngestError("the PDF could not be read") from exc
    text = "\n\n".join(pages)
    if not text:
        raise PermanentIngestError(
            "the PDF has no text layer (a scanned PDF?); only text PDFs are supported"
        )
    return text


def _docx(raw: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            sizes = [info.file_size for info in archive.infolist()]
        if sizes and max(max(sizes), sum(sizes)) > MAX_INFLATED_BYTES:
            raise PermanentIngestError(TOO_LONG)
        doc = DocxDocument(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
    except PermanentIngestError:
        raise
    except Exception as exc:  # noqa: BLE001 — python-docx raises many types; all permanent
        raise PermanentIngestError("the Word document could not be read") from exc
    return "\n\n".join(parts)

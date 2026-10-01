"""Upload bytes -> text for each supported format; unreadable input is a
PermanentIngestError with a user-readable message."""

import io

import pytest
from docx import Document as DocxDocument

from packages.core.rag.errors import PermanentIngestError
from packages.core.rag.extract import extract_text


def _pdf(text: str | None) -> bytes:
    """A one-page PDF; text=None gives a page with no text layer."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
        b" /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objs) + 1, xref)
    )
    return out.getvalue()


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_text_formats_decode_utf8():
    assert extract_text("héllo".encode(), "notes.md") == "héllo"


def test_bad_utf8_is_permanent():
    with pytest.raises(PermanentIngestError):
        extract_text(b"\xff\xfe\xfa", "notes.txt")


def test_pdf_text_layer():
    assert "Quarterly revenue grew" in extract_text(
        _pdf("Quarterly revenue grew"), "Report.PDF"
    )


def test_pdf_without_text_is_permanent_with_a_reason():
    with pytest.raises(PermanentIngestError, match="no text layer"):
        extract_text(_pdf(None), "scan.pdf")


def test_garbage_pdf_is_permanent():
    with pytest.raises(PermanentIngestError, match="could not be read"):
        extract_text(b"not a pdf", "x.pdf")


def test_docx_paragraphs_and_tables():
    text = extract_text(
        _docx(["First paragraph.", "Second."], [["Ticker", "Weight"], ["ABC", "5%"]]),
        "memo.docx",
    )
    assert "First paragraph." in text and "Second." in text
    assert "Ticker | Weight" in text and "ABC | 5%" in text


def test_garbage_docx_is_permanent():
    with pytest.raises(PermanentIngestError, match="could not be read"):
        extract_text(b"PK not really a zip", "x.docx")

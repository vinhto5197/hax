"""The parse itself is bounded: a PDF past MAX_PDF_PAGES pages and a DOCX past
MAX_DOCX_PARAGRAPHS paragraphs are a PermanentIngestError, so a huge file fails
fast instead of being parsed to the end. A file exactly at the budget is still
read in full. The refusal is user-visible and static, so it never echoes the
document's content. The real constants are used, not monkeypatched, so the
budgets the worker ships with are the ones pinned here."""

import io

import pytest
from docx import Document as DocxDocument
from pypdf import PdfReader, PdfWriter

from packages.core.rag.errors import PermanentIngestError
from packages.core.rag.extract import MAX_DOCX_PARAGRAPHS, MAX_PDF_PAGES, extract_text

MARKER = "zebraquartz"


def _text_page_pdf(text: str) -> bytes:
    """A one-page PDF whose page carries a Helvetica text layer."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
        b" /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objs) + 1, xref)
    )
    return out.getvalue()


def _pdf(pages: int, text: str | None = None) -> bytes:
    """A PDF of `pages` pages; with `text`, the first page carries it and the
    rest are blank."""
    writer = PdfWriter()
    if text is not None:
        writer.append(PdfReader(io.BytesIO(_text_page_pdf(text))))
    while len(writer.pages) < pages:
        writer.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    writer.write(out)
    raw = out.getvalue()
    assert len(PdfReader(io.BytesIO(raw)).pages) == pages
    return raw


def _docx(paragraphs: int) -> bytes:
    """A DOCX of exactly `paragraphs` body paragraphs, each carrying MARKER;
    counted on the document so the template's own paragraphs are included."""
    doc = DocxDocument()
    # Counted once: doc.paragraphs re-queries the XML on every access.
    for i in range(len(doc.paragraphs), paragraphs):
        doc.add_paragraph(f"{MARKER} {i}")
    assert len(doc.paragraphs) == paragraphs
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def docx_over_budget() -> bytes:
    return _docx(MAX_DOCX_PARAGRAPHS + 1)


@pytest.fixture(scope="module")
def docx_at_budget() -> bytes:
    return _docx(MAX_DOCX_PARAGRAPHS)


def test_budgets_are_the_shipped_values():
    assert MAX_PDF_PAGES == 300
    assert MAX_DOCX_PARAGRAPHS == 5000


def test_pdf_past_the_page_budget_is_permanent():
    with pytest.raises(PermanentIngestError) as excinfo:
        extract_text(_pdf(MAX_PDF_PAGES + 1), "big.pdf")

    assert "too many pages" in str(excinfo.value).lower()


def test_pdf_at_the_page_budget_is_read():
    text = extract_text(_pdf(MAX_PDF_PAGES, text=f"{MARKER} revenue grew"), "a.pdf")

    assert MARKER in text


def test_docx_past_the_paragraph_budget_is_permanent_and_does_not_echo_content(
    docx_over_budget,
):
    with pytest.raises(PermanentIngestError) as excinfo:
        extract_text(docx_over_budget, "big.docx")

    message = str(excinfo.value)
    assert "too many paragraphs" in message.lower()
    assert MARKER not in message


def test_docx_at_the_paragraph_budget_is_read(docx_at_budget):
    text = extract_text(docx_at_budget, "memo.docx")

    assert MARKER in text
    assert f"{MARKER} {MAX_DOCX_PARAGRAPHS - 1}" in text
